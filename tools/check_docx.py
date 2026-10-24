# -*- coding: utf-8 -*-
"""验证 docx 内容与结构，报告写入 UTF-8 文件"""
import io, sys
from docx import Document

path = r"C:\Users\32375\Desktop\51单片机智能门锁代码精讲.docx"
out = r"E:\c51_project\智能门锁\tools\check_report.txt"
d = Document(path)

lines = []
lines.append(f"paragraphs: {len(d.paragraphs)}  tables: {len(d.tables)}")

h1 = [p.text for p in d.paragraphs if p.runs and p.runs[0].font.size and p.runs[0].font.size.pt == 22]
lines.append(f"H1(22pt) count: {len(h1)} -> {h1}")

# 各级标题统计
sizes = {}
for p in d.paragraphs:
    if p.runs and p.runs[0].font.size and p.runs[0].font.bold:
        s = p.runs[0].font.size.pt
        sizes[s] = sizes.get(s, 0) + 1
lines.append(f"heading size histogram: {sizes}")

# 章节标题
chapter_heads = [p.text for p in d.paragraphs
                 if p.runs and p.runs[0].font.size
                 and p.runs[0].font.size.pt == 16
                 and (p.text.startswith("第") or p.text.startswith("附录"))]
lines.append(f"chapter headings ({len(chapter_heads)}):")
for h in chapter_heads:
    lines.append("   " + h)

# 检查代码块样式（8.5pt 灰底）
code_blocks = [p for p in d.paragraphs
               if p.runs and p.runs[0].font.size and p.runs[0].font.size.pt == 8.5]
lines.append(f"code blocks: {len(code_blocks)}")

# 表格行/列
for i, t in enumerate(d.tables[:5]):
    lines.append(f"table{i}: rows={len(t.rows)} cols={len(t.columns)} first={t.cell(0,0).text[:24]}")

# 检查是否有空段落堆积
empty_runs = sum(1 for p in d.paragraphs if not p.text.strip())
lines.append(f"empty paragraphs: {empty_runs}")

# 检查关键内容是否齐全
keywords = ["目录", "第 0 章", "第 1 章", "第 2 章", "第 3 章", "第 4 章", "第 5 章",
            "第 6 章", "第 7 章", "第 8 章", "第 9 章", "附录 A", "附录 B", "附录 C", "附录 D",
            "Bug 6", "Bug 10", "结束语"]
full = "\n".join(p.text for p in d.paragraphs)
missing = [k for k in keywords if k not in full]
lines.append(f"missing keywords: {missing if missing else 'none'}")

with io.open(out, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("report written:", out)
